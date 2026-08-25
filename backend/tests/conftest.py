"""Shared fixtures. All sample files are built in-code (openpyxl/pymupdf/zipfile) —
no binary fixtures are committed, so nothing resembling a real client document
ever sits in the repo.
"""
import io
import os
import sys
import zipfile

import openpyxl
import pymupdf
import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from main import app, sessions  # noqa: E402


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture(autouse=True)
def clear_sessions():
    """Sessions are a module-level dict — reset between tests so they don't leak."""
    sessions.clear()
    yield
    sessions.clear()


def make_excel_bytes(headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def excel_bytes():
    return make_excel_bytes(
        ["Name", "Date", "Amount"],
        [
            ["John Doe", "2024-01-15", 5000],
            ["Jane Smith", "2024-02-20", 7500],
        ],
    )


def make_pdf_bytes(placeholders):
    doc = pymupdf.open()
    page = doc.new_page()
    for i, ph in enumerate(placeholders):
        page.insert_text((72, 72 + i * 20), f"<<{ph}>>")
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


@pytest.fixture
def pdf_bytes():
    return make_pdf_bytes(["Name", "Date"])


def make_docx_bytes(field_names):
    fields_xml = "".join(
        '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r><w:instrText xml:space="preserve"> MERGEFIELD {name} \\* MERGEFORMAT </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        f'<w:r><w:t>«{name}»</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
        for name in field_names
    )
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{fields_xml}</w:body></w:document>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", document_xml)
    return buf.getvalue()


@pytest.fixture
def docx_bytes():
    return make_docx_bytes(["Name", "Date"])


_CONTENT_TYPES_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    "</Types>"
)

_ROOT_RELS_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
    'Target="word/document.xml"/>'
    "</Relationships>"
)


def make_valid_docx_bytes(field_names):
    """A complete, minimal OOXML package real Word-processing software (LibreOffice
    included) can actually open — unlike make_docx_bytes above, which writes only
    word/document.xml and is sufficient for this app's own zipfile-based parsing
    but is not a valid docx package on its own."""
    fields_xml = "".join(
        '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r><w:instrText xml:space="preserve"> MERGEFIELD {name} \\* MERGEFORMAT </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        f'<w:r><w:t>«{name}»</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
        for name in field_names
    )
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{fields_xml}</w:body></w:document>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", _CONTENT_TYPES_XML)
        z.writestr("_rels/.rels", _ROOT_RELS_XML)
        z.writestr("word/document.xml", document_xml)
    return buf.getvalue()
